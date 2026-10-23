#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import html
import json
import math
import os
import re
import sys
import tempfile
import urllib.error
import urllib.request
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

import duckdb
from PIL import Image, ImageOps
import openpyxl
from shapely import transform as transform_geometry
from shapely.geometry import shape
from shapely.ops import unary_union

CITY_BOUNDARIES_PATH = Path(__file__).resolve().parents[1] / "src" / "data" / "san-fernando-barangays.geojson"
MAP_TILE_ZOOM = 13

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "src" / "data"
DEFAULT_CACHE_DIR = ROOT / "scripts" / ".cache"
API_BASE = "https://data.bettergov.ph/api/v1"
NEWS_API_BASE = "https://www.sanfernandocity.gov.ph/wp-json/wp/v2"
NEWS_STORY_LIMIT = 30
NEWS_IMAGE_SIZE = (1200, 800)
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
)
BLGF_REFERER = "https://blgf.gov.ph/lgu-fiscal-data/"

SOURCE_SPECS: dict[str, dict[str, Any]] = {
    "sre": {
        "dataset_id": 9,
        "filename": "blgf-sre-2024.xlsx",
        "terms": ("fy 2024 statement of receipts and expenditures by lgu",),
        "fallback_url": "https://blgf.gov.ph/wp-content/uploads/2025/08/By-LGU-SRE-2024.xlsx",
        "referer": BLGF_REFERER,
    },
    "ari": {
        "dataset_id": 9,
        "filename": "blgf-ari-2024.xlsx",
        "terms": ("fy 2024 annual regular income and dependencies by lgu",),
        "fallback_url": "https://blgf.gov.ph/wp-content/uploads/2025/06/By-LGU-ARI-and-Dependencies-2024.xlsx",
        "referer": BLGF_REFERER,
    },
    "ldrrmf": {
        "dataset_id": 9,
        "filename": "blgf-ldrrmf-2024.xlsx",
        "terms": ("fy 2024 local disaster risk reduction management fund by lgu",),
        "fallback_url": "https://blgf.gov.ph/wp-content/uploads/2025/06/FY2024-LDRRMF-by-LGU.xlsx",
        "referer": BLGF_REFERER,
    },
    "sef": {
        "dataset_id": 9,
        "filename": "blgf-sef-2024.xlsx",
        "terms": ("fy 2024 special education fund income and expenditures",),
        "fallback_url": "https://blgf.gov.ph/wp-content/uploads/2025/06/FY2024-SEF-Income-and-Expenditures.xlsx",
        "referer": BLGF_REFERER,
    },
    "elections": {
        "dataset_id": 25,
        "filename": "halalan-winners.csv",
        "terms": ("election winners", "2001-2025"),
        "fallback_url": "https://github.com/RobertRLeung/OpenHalalan/releases/download/data-latest/NLE_Winners_2004-2025.csv",
    },
    "boundaries": {
        "dataset_id": 23,
        "filename": "barangays.geojson",
        "terms": ("barangays", "geojson"),
        "fallback_url": "https://github.com/bendlikeabamboo/barangay-boundaries-repository/releases/download/v2026.4.13.0/barangays.geojson",
    },
    "dpwh": {
        "dataset_id": 19,
        "filename": "dpwh.parquet",
        "terms": ("dpwh infrastructure projects data",),
        "fallback_url": "https://huggingface.co/datasets/bettergovph/dpwh-transparency-data/resolve/main/dpwh_transparency_data.parquet?download=true",
    },
    "philgeps": {
        "dataset_id": 5,
        "filename": "philgeps.parquet",
        "terms": ("philgeps complete dataset",),
        "fallback_url": "https://huggingface.co/datasets/bettergovph/philgeps-data/resolve/main/philgeps.parquet?download=true",
    },
}

PARTY_LABELS = {
    "IND": "Independent",
    "PDPLBN": "PDP-Laban",
    "LAKAS": "Lakas",
    "LAKAS NUCD-UMDP": "Lakas-NUCD-UMDP",
    "LAKAS-CMD": "Lakas-CMD",
    "NPC": "NPC",
    "NP": "NP",
    "KANP": "KANP",
}

NAME_OVERRIDES = {
    "ORTEGA, MARY JANE C.": "Mary Jane C. Ortega",
    "ORTEGA, MARY JANE CRISANTO": "Mary Jane C. Ortega",
    "VALERO, BERNARDO LUCERO": "Bernardo L. Valero",
    "NISCE, PANCRACIO QUEROL": "Pancracio Q. Nisce",
    "GUALBERTO, HERMENEGILDO ANDAYA": "Hermenegildo 'Dong' A. Gualberto",
    "GUALBERTO, DONG ANDAYA": "Hermenegildo 'Dong' A. Gualberto",
    "GUALBERTO, DONG": "Hermenegildo 'Dong' A. Gualberto",
    "ORTEGA, PABLO CAMPOS": "Pablo C. Ortega",
}

ELECTION_SUPPLEMENTS = {
    2007: {
        "sourceUrl": (
            "https://www.comelec.gov.ph/php-tpls-attachments/"
            "ListElectedCandidates/2007_list_of_elected_city_municipal_candidates.pdf"
        ),
        "winners": {
            "mayor": {"name": "Pablo C. Ortega", "party": "Lakas-CMD"},
        },
    },
    2010: {
        "sourceUrl": (
            "https://www.comelec.gov.ph/php-tpls-attachments/"
            "ListElectedCandidates/2010_list_of_elected_citymunicipal_candidates.pdf"
        ),
        "winners": {
            "mayor": {
                "name": "Pablo C. Ortega",
                "party": "Lakas-Kampi-CMD (LKS-KAM)",
            },
            "viceMayor": {
                "name": "Hermenegildo A. Gualberto",
                "party": "Independent",
            },
        },
    },
    2013: {
        "sourceUrl": (
            "https://www.comelec.gov.ph/php-tpls-attachments/"
            "ListElectedCandidates/2013_list_of_elected_citymunicipal_candidates.pdf"
        ),
        "winners": {
            "mayor": {"name": "Pablo C. Ortega", "party": "Independent"},
            "viceMayor": {
                "name": "Hermenegildo A. Gualberto",
                "party": "Independent",
            },
        },
    },
}

PROJECT_TITLE_OVERRIDES = {
    "24AE0103": "Slope protection structure",
    "25AE0070": "Multi-purpose facilities and site development",
    "24AE0088": "Seawall, Barangay Pagdaraoan",
    "24AE0090": "Seawall, Barangay Ilocanos Sur",
    "24AE0089": "Seawall, Barangay Ilocanos Norte",
}

CONTRACTOR_OVERRIDES = {
    "25AE0070": "Octagon Concrete Solutions",
}

AWARDEE_OVERRIDES = {
    "LAV CONSTRUCTION & SUPPLY": "LAV Construction & Supply",
}



def normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def place_key(value: Any) -> str:
    text = str(value).lower().replace("(pob.)", "").replace("poblacion", "")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def quote(value: Any) -> str:
    text = str(value)
    if "'" in text and '"' not in text:
        return json.dumps(text, ensure_ascii=False)
    if '"' not in text:
        escaped = (
            text.replace("\\", "\\\\")
            .replace("'", "\\'")
            .replace("\n", "\\n")
            .replace("\r", "\\r")
            .replace("\t", "\\t")
        )
        return f"'{escaped}'"
    return json.dumps(text, ensure_ascii=False)


def num(value: Any, places: int = 2) -> str:
    rounded = round(float(value), places)
    if rounded == int(rounded):
        return str(int(rounded))
    return f"{rounded:.{places}f}".rstrip("0").rstrip(".")


def as_number(value: Any, label: str) -> float:
    if value is None or value == "":
        raise ValueError(f"{label} is empty")
    if isinstance(value, str):
        value = value.replace(",", "").strip()
    return float(value)


def find_row(rows: list[tuple[Any, ...]], predicate: Any, label: str) -> tuple[Any, ...]:
    for row in rows:
        if len(row) > 4 and predicate(row):
            return row
    raise ValueError(f"Could not find {label}")


def read_sheet_rows(path: Path, sheet_contains: str) -> list[tuple[Any, ...]]:
    workbook = openpyxl.load_workbook(filename=path, read_only=True, data_only=True)
    try:
        matches = [name for name in workbook.sheetnames if sheet_contains.lower() in name.lower()]
        if not matches:
            raise ValueError(f"No sheet matching {sheet_contains!r} in {path.name}")
        return list(workbook[matches[0]].iter_rows(values_only=True))
    finally:
        workbook.close()


def read_json(url: str) -> Any:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": USER_AGENT, "Accept": "application/json"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.load(response)


def api_json(path: str) -> Any:
    return read_json(f"{API_BASE}{path}")


def resolve_sources(offline: bool) -> dict[str, dict[str, Any]]:
    sources: dict[str, dict[str, Any]] = {}
    for key, spec in SOURCE_SPECS.items():
        url = spec["fallback_url"]
        if not offline:
            try:
                payload = api_json(
                    f"/resources?dataset_id={spec['dataset_id']}&limit=100"
                )
                resources = payload.get("data", []) if isinstance(payload, dict) else payload
                match = next(
                    (
                        resource
                        for resource in resources
                        if all(
                            normalize(term) in normalize(resource.get("name", ""))
                            for term in spec["terms"]
                        )
                    ),
                    None,
                )
                if match and match.get("download_url"):
                    url = match["download_url"]
                    print(f"resolved {key}: {match['name']}")
                else:
                    print(f"warning: resource lookup missed {key}; using pinned URL")
            except Exception as exc:
                print(f"warning: resource lookup failed for {key}: {exc}")
        sources[key] = {**spec, "url": url}
    return sources


def download_source(
    key: str,
    spec: dict[str, Any],
    cache_dir: Path,
    force: bool,
    offline: bool,
    cached_fallbacks: list[str],
) -> Path:
    destination = cache_dir / spec["filename"]
    if offline:
        if not destination.exists():
            raise FileNotFoundError(f"Missing cached source for {key}: {destination}")
        print(f"using cached {key}: {destination.name}")
        return destination

    url = spec["url"]
    meta_path = destination.with_name(f"{destination.name}.meta.json")
    meta: dict[str, Any] = {}
    if meta_path.exists():
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            meta = {}

    headers = {"User-Agent": USER_AGENT, "Accept": "*/*"}
    if spec.get("referer"):
        headers["Referer"] = spec["referer"]
    if destination.exists() and not force and meta.get("url") == url:
        if meta.get("etag"):
            headers["If-None-Match"] = meta["etag"]
        if meta.get("last_modified"):
            headers["If-Modified-Since"] = meta["last_modified"]

    request = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            content_length = response.headers.get("Content-Length")
            total = int(content_length) if content_length and content_length.isdigit() else 0
            print(
                f"downloading {key}: {spec['filename']}"
                + (f" ({total / (1024 * 1024):.1f} MB)" if total else "")
            )
            temp_path: Path | None = None
            written = 0
            next_report = 0
            try:
                with tempfile.NamedTemporaryFile(
                    mode="wb",
                    dir=cache_dir,
                    prefix=f"{spec['filename']}.",
                    suffix=".part",
                    delete=False,
                ) as handle:
                    temp_path = Path(handle.name)
                    while True:
                        chunk = response.read(1024 * 1024)
                        if not chunk:
                            break
                        handle.write(chunk)
                        written += len(chunk)
                        if total and written >= next_report:
                            print(f"  {written * 100 / total:.0f}%", flush=True)
                            next_report += max(total // 5, 1)
                os.replace(temp_path, destination)
            finally:
                if temp_path and temp_path.exists():
                    temp_path.unlink()
    except urllib.error.HTTPError as exc:
        if exc.code == 304 and destination.exists():
            print(f"cached {key}: unchanged")
            return destination
        if destination.exists() and not force:
            print(f"warning: download failed for {key}; using cached copy")
            cached_fallbacks.append(key)
            return destination
        raise
    except Exception as exc:
        if destination.exists() and not force:
            print(f"warning: download failed for {key}: {exc}; using cached copy")
            cached_fallbacks.append(key)
            return destination
        raise

    try:
        meta_path.write_text(
            json.dumps(
                {
                    "url": url,
                    "etag": response.headers.get("ETag"),
                    "last_modified": response.headers.get("Last-Modified"),
                },
                indent=2,
            ),
            encoding="utf-8",
        )
    except Exception:
        pass
    return destination


def clean_html(value: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", value))).strip()


def news_summary(content: str, title: str) -> str:
    paragraphs = [clean_html(part) for part in re.findall(r"<p[^>]*>(.*?)</p>", content, re.S)]
    paragraphs = [part for part in paragraphs if part]
    normalized_title = re.sub(r"[’‘]", "'", title).casefold()
    body = next(
        (
            part
            for part in paragraphs
            if re.sub(r"[’‘]", "'", part).casefold() != normalized_title
        ),
        "",
    )
    if not body:
        body = clean_html(content)
    if re.sub(r"[’‘]", "'", body).casefold().startswith(normalized_title):
        body = body[len(title) :].lstrip(" -–—:.")
    return body.strip()


def fetch_news_image(url: str) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": USER_AGENT, "Accept": "image/*"},
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        with Image.open(BytesIO(response.read())) as source:
            image = ImageOps.exif_transpose(source).convert("RGB")
            image = ImageOps.fit(image, NEWS_IMAGE_SIZE, method=Image.Resampling.LANCZOS)
            output = BytesIO()
            image.save(output, format="JPEG", quality=82, optimize=True, progressive=True)
            return output.getvalue()


def fetch_news() -> list[tuple[dict[str, Any], bytes]]:
    fields = "date,slug,title,content,link,featured_media"
    url = (
        f"{NEWS_API_BASE}/posts?per_page={NEWS_STORY_LIMIT}"
        f"&orderby=date&order=desc&_fields={fields}"
    )
    posts = read_json(url)
    if not isinstance(posts, list) or not posts:
        raise ValueError("WordPress returned no news posts")

    stories: list[tuple[dict[str, Any], bytes]] = []
    for post in posts[:NEWS_STORY_LIMIT]:
        title = clean_html(post["title"]["rendered"])
        slug = str(post["slug"])
        if not re.fullmatch(r"[a-z0-9-]+", slug):
            raise ValueError(f"Unexpected news slug: {slug}")
        excerpt = news_summary(str(post["content"]["rendered"]), title)
        if not excerpt:
            raise ValueError(f"No summary found for news story: {slug}")
        media_id = post.get("featured_media")
        if not media_id:
            raise ValueError(f"No featured image for news story: {slug}")
        media = read_json(f"{NEWS_API_BASE}/media/{media_id}?_fields=source_url")
        image_url = str(media.get("source_url") or "")
        if not image_url.startswith("https://"):
            raise ValueError(f"No image URL for news story: {slug}")
        stories.append(
            (
                {
                    "slug": slug,
                    "title": title,
                    "date": str(post["date"]),
                    "excerpt": excerpt,
                    "image": f"/news/news-{len(stories) + 1}.jpg",
                    "sourceUrl": str(post["link"]),
                },
                fetch_news_image(image_url),
            )
        )
    return stories


def write_news(stories: list[tuple[dict[str, Any], bytes]]) -> None:
    news_dir = ROOT / "public" / "news"
    news_dir.mkdir(parents=True, exist_ok=True)
    for index, (_, image) in enumerate(stories, start=1):
        write_bytes(news_dir / f"news-{index}.jpg", image)
    for path in news_dir.glob("news-*.jpg"):
        if int(path.stem.split("-")[1]) > len(stories):
            path.unlink()

    lines = [
        "import type { NewsItem } from './types';",
        "",
        f"export const newsLastVerified = {json.dumps(date.today().isoformat())}",
        "",
        "export const news: NewsItem[] = [",
    ]
    for item, _ in stories:
        lines.extend(
            [
                "  {",
                f"    \"slug\": {json.dumps(item['slug'], ensure_ascii=False)},",
                f"    \"title\": {json.dumps(item['title'], ensure_ascii=False)},",
                f"    \"date\": {json.dumps(item['date'], ensure_ascii=False)},",
                f"    \"excerpt\": {json.dumps(item['excerpt'], ensure_ascii=False)},",
                f"    \"image\": {json.dumps(item['image'], ensure_ascii=False)},",
                f"    \"sourceUrl\": {json.dumps(item['sourceUrl'], ensure_ascii=False)}",
                "  },",
            ]
        )
    lines.extend(["]", ""])
    write_text(DATA_DIR / "news.ts", "\n".join(lines))


def refresh_news(offline: bool, cached_fallbacks: list[str]) -> int | None:
    if offline:
        print("news: offline; keeping existing snapshot")
        return None
    try:
        stories = fetch_news()
        write_news(stories)
    except Exception as exc:
        print(f"warning: news refresh failed: {exc}; keeping existing snapshot")
        cached_fallbacks.append("news")
        return None
    print(f"refreshed {len(stories)} news stories")
    return len(stories)


def parse_fiscal(
    sre_path: Path, ari_path: Path, ldrrmf_path: Path, sef_path: Path
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    sre_rows = read_sheet_rows(sre_path, "SRE")
    ari_rows = read_sheet_rows(ari_path, "ARI")
    ldrrmf_rows = read_sheet_rows(ldrrmf_path, "FY2024")
    sef_rows = read_sheet_rows(sef_path, "Income")

    def is_sf(row: tuple[Any, ...], province_col: int, name_col: int) -> bool:
        return (
            normalize(row[1]) == "region i"
            and normalize(row[province_col]) == "la union"
            and "san fernando" in normalize(row[name_col])
        )

    sre = find_row(
        sre_rows,
        lambda row: len(row) > 43 and is_sf(row, 2, 3),
        "San Fernando City in the FY2024 SRE",
    )
    ldrrmf = find_row(
        ldrrmf_rows,
        lambda row: len(row) > 10
        and normalize(row[1]) == "city"
        and normalize(row[2]) == "region i"
        and normalize(row[3]) == "la union"
        and "san fernando" in normalize(row[4]),
        "San Fernando City in the FY2024 LDRRMF",
    )
    sef = find_row(
        sef_rows,
        lambda row: len(row) > 5 and is_sf(row, 2, 3),
        "San Fernando City in the FY2024 SEF",
    )

    fiscal = {
        "totalIncome": as_number(sre[22], "SRE total income"),
        "totalLocalSources": as_number(sre[16], "SRE local sources"),
        "totalTaxRevenue": as_number(sre[10], "SRE tax revenue"),
        "totalNonTaxRevenue": as_number(sre[15], "SRE non-tax revenue"),
        "nta": as_number(sre[17], "SRE NTA"),
        "otherNationalShares": as_number(sre[18], "SRE other national shares"),
        "grants": as_number(sre[20], "SRE grants"),
        "totalOperatingExpenditures": as_number(sre[32], "SRE operating expenditures"),
        "netOperatingIncome": as_number(sre[33], "SRE net operating income"),
        "capitalOutlay": as_number(sre[43], "SRE capital outlay"),
        "specialEducationFund": as_number(sef[5], "SEF income"),
        "disasterFund": {
            "appropriation": as_number(ldrrmf[9], "LDRRMF appropriation"),
            "spent": as_number(ldrrmf[10], "LDRRMF spent"),
            "quickResponseAppropriation": as_number(
                ldrrmf[7], "LDRRMF quick response appropriation"
            ),
            "quickResponseSpent": as_number(ldrrmf[8], "LDRRMF quick response spent"),
        },
    }

    peers: list[dict[str, Any]] = []
    for row in ari_rows:
        if (
            len(row) > 25
            and normalize(row[1]) == "region i"
            and normalize(row[4]) == "city"
            and row[3]
        ):
            name = re.sub(r"\s+City$", "", str(row[3]).strip(), flags=re.I)
            peers.append(
                {
                    "name": name,
                    "province": str(row[2]).strip(),
                    "annualRegularIncome": as_number(row[23], f"{name} ARI"),
                    "localSources": as_number(row[13], f"{name} local sources"),
                    "ownSourceShare": as_number(row[24], f"{name} own-source share"),
                    "ntaShare": as_number(row[25], f"{name} NTA share"),
                }
            )

    if not any(peer["name"] == "San Fernando" for peer in peers):
        raise ValueError("San Fernando City is missing from the FY2024 ARI cities")
    return fiscal, peers


def display_person(raw: str) -> str:
    value = re.sub(r"\s+", " ", raw.strip().upper())
    if value in NAME_OVERRIDES:
        return NAME_OVERRIDES[value]
    if "," in value:
        last, given = value.split(",", 1)
        given = " ".join(
            token[:1].upper() + token[1:].lower() if len(token) > 2 else token
            for token in given.strip().split()
        )
        last = " ".join(token.capitalize() for token in last.strip().split())
        return f"{given} {last}"
    return value.title()


def display_party(raw: str) -> str:
    value = re.sub(r"\s+", " ", raw.strip().upper())
    return PARTY_LABELS.get(value, value.title())


def parse_elections(path: Path) -> list[dict[str, Any]]:
    elections: dict[int, dict[str, Any]] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if normalize(row.get("Province", "")) != "la union":
                continue
            if "san fernando" not in normalize(row.get("City", "")):
                continue
            position = normalize(row.get("Position", ""))
            if position not in {"mayor", "vice mayor"}:
                continue
            try:
                year = int(row["Year"])
            except (KeyError, TypeError, ValueError):
                continue
            elections.setdefault(year, {"year": year, "mayor": None, "viceMayor": None})
            key = "mayor" if position == "mayor" else "viceMayor"
            elections[year][key] = {
                "name": display_person(row["Full Name"]),
                "party": display_party(row["Party"]),
            }
    for year, supplement in ELECTION_SUPPLEMENTS.items():
        election = elections.setdefault(
            year, {"year": year, "mayor": None, "viceMayor": None}
        )
        for position, winner in supplement["winners"].items():
            if election.get(position) is None:
                election[position] = winner
    return [elections[year] for year in sorted(elections)]


def display_contractor(raw: str) -> str:
    text = re.sub(r"\s+", " ", str(raw or "").strip())
    text = re.sub(r"\s*\(\s*\[?REVOKED\]?\s*\d+\s*\)\s*$", "", text, flags=re.I)
    text = re.sub(r"\s*\(\d+\)\s*$", "", text)
    return text.strip().title()


def display_project_title(contract_id: str, raw: str) -> str:
    if contract_id in PROJECT_TITLE_OVERRIDES:
        return PROJECT_TITLE_OVERRIDES[contract_id]
    text = re.sub(r"\s+", " ", str(raw or "").strip())
    text = re.sub(
        r"^(SUSTAINABLE INFRASTRUCTURE PROJECTS ALLEVIATING GAPS \(SIPAG\)|"
        r"BASIC INFRASTRUCTURE PROGRAM \(BIP\))\s*-\s*",
        "",
        text,
        flags=re.I,
    )
    text = re.sub(
        r",?\s*(CITY OF\s+)?SAN FERNANDO( CITY)?(?:,\s*LA UNION)?.*$",
        "",
        text,
        flags=re.I,
    )
    text = re.sub(r"\s*\(RE-ADVERTISEMENT\)\s*$", "", text, flags=re.I)
    text = text.replace("BRGY.", "Barangay").strip(" ,-")
    if text.isupper():
        text = text.lower()
    return text[:1].upper() + text[1:]


def display_status(raw: str) -> str:
    return {
        "on going": "Ongoing",
        "completed": "Completed",
        "for procurement": "For procurement",
    }.get(normalize(raw), str(raw).strip())


def compact_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    return text


def optional_number(value: Any, places: int = 6) -> float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(number):
        return None
    return round(number, places)


def load_barangay_names() -> list[str]:
    source_text = (DATA_DIR / "barangays.ts").read_text(encoding="utf-8")
    declaration = source_text.index("export const barangays")
    array_start = source_text.index("[", source_text.index("= [", declaration))
    array_end = source_text.rindex("]")
    sites = json.loads(source_text[array_start : array_end + 1])
    return [site["name"] for site in sites]


BARANGAY_NAMES = load_barangay_names()


def find_project_barangay(description: Any, barangay_names: list[str]) -> str | None:
    keys = [(compact_key(name), name) for name in barangay_names]
    for match in re.finditer(r"(?:BARANGAY|BRGY\.?)\s*([A-Z0-9 .'-]+)", str(description).upper()):
        tail = compact_key(match.group(1))
        candidates = [(len(key), name) for key, name in keys if key and key in tail]
        if candidates:
            return max(candidates)[1]
    return None


def project_category(raw: Any, component: Any) -> str:
    value = str(component or "").strip() or str(raw or "").strip()
    allowed = {"Buildings and Facilities", "Roads", "Flood Control and Drainage", "Water Provision and Storage", "Bridges"}
    return value if value in allowed else "Other"


def parse_dpwh(path: Path) -> dict[str, Any]:
    connection = duckdb.connect()
    try:
        connection.execute(
            "CREATE TABLE dpwh AS SELECT * FROM read_parquet(?)", [str(path)]
        )
        where = (
            "lower(location.province) LIKE '%la union%' "
            "AND lower(description) LIKE '%san fernando%'"
        )
        total_projects, total_budget = connection.execute(
            f"SELECT count(*), coalesce(sum(budget), 0) FROM dpwh WHERE {where}"
        ).fetchone()
        status_counts = dict(
            connection.execute(
                f"SELECT lower(status), count(*) FROM dpwh WHERE {where} GROUP BY 1"
            ).fetchall()
        )
        recent_years = connection.execute(
            f"""
            SELECT TRY_CAST(infraYear AS INTEGER) AS year,
                   count(*) AS projects,
                   coalesce(sum(budget), 0) AS amount
            FROM dpwh
            WHERE {where}
              AND TRY_CAST(infraYear AS INTEGER) BETWEEN 2021 AND 2025
            GROUP BY 1
            ORDER BY 1
            """
        ).fetchall()
        largest = connection.execute(
            f"""
            SELECT contractId, description, status, budget, progress, contractor,
                   TRY_CAST(infraYear AS INTEGER) AS year
            FROM dpwh
            WHERE {where}
            ORDER BY budget DESC
            LIMIT 5
            """
        ).fetchall()
        records = connection.execute(
            f"""
            SELECT contractId, description, status, budget, progress, contractor,
                   latitude, longitude, startDate, completionDate,
                   TRY_CAST(infraYear AS INTEGER) AS year, category,
                   componentCategories, sourceOfFunds
            FROM dpwh
            WHERE {where}
            ORDER BY budget DESC
            """
        ).fetchall()
    finally:
        connection.close()

    projects = []
    for contract_id, description, status, budget, progress, contractor, year in largest:
        contract_id = str(contract_id or "")
        projects.append(
            {
                "contractId": contract_id,
                "title": display_project_title(contract_id, description),
                "year": int(year),
                "budget": as_number(budget, f"{contract_id} budget"),
                "status": display_status(status),
                "progress": as_number(progress, f"{contract_id} progress"),
                "contractor": CONTRACTOR_OVERRIDES.get(
                    contract_id, display_contractor(contractor)
                ),
            }
        )

    if not projects:
        raise ValueError("No DPWH projects matched San Fernando")
    return {
        "totalProjects": int(total_projects),
        "totalBudget": as_number(total_budget, "DPWH total budget"),
        "completed": int(status_counts.get("completed", 0)),
        "ongoing": int(status_counts.get("on-going", 0)),
        "forProcurement": int(status_counts.get("for procurement", 0)),
        "recentYears": [
            {
                "year": int(year),
                "projects": int(project_count),
                "amount": as_number(amount, f"{year} DPWH amount"),
            }
            for year, project_count, amount in recent_years
        ],
        "largestProjects": projects,
        "records": [
            {
                "contractId": str(contract_id or ""),
                "title": display_project_title(str(contract_id or ""), description),
                "year": int(year or 0),
                "budget": round(as_number(budget, f"{contract_id} budget"), 2),
                "status": display_status(status),
                "progress": round(as_number(progress, f"{contract_id} progress"), 2),
                "contractor": CONTRACTOR_OVERRIDES.get(
                    str(contract_id or ""), display_contractor(contractor)
                ),
                "category": project_category(category, component_categories),
                "barangay": find_project_barangay(description, BARANGAY_NAMES),
                "latitude": optional_number(latitude),
                "longitude": optional_number(longitude),
                "startDate": optional_text(start_date),
                "completionDate": optional_text(completion_date),
                "sourceOfFunds": optional_text(source_of_funds),
            }
            for (
                contract_id,
                description,
                status,
                budget,
                progress,
                contractor,
                latitude,
                longitude,
                start_date,
                completion_date,
                year,
                category,
                component_categories,
                source_of_funds,
            ) in records
        ],
    }


def display_awardee(raw: str) -> str:
    value = re.sub(r"\s+", " ", str(raw or "").strip().upper())
    return AWARDEE_OVERRIDES.get(value, value.title())


def display_award_title(raw: str) -> str:
    text = re.sub(r"\s+", " ", str(raw or "").strip())
    text = re.sub(r"\bRehab\.", "Rehabilitation", text, flags=re.I)
    if text.isupper():
        text = text.lower()
    text = text[:1].upper() + text[1:]
    return text


def parse_philgeps(path: Path) -> dict[str, Any]:
    connection = duckdb.connect()
    try:
        connection.execute(
            "CREATE TABLE philgeps AS SELECT * FROM read_parquet(?)", [str(path)]
        )
        where = (
            "upper(trim(organization_name)) = 'CITY OF SAN FERNANDO, LA UNION' "
            "AND award_date >= DATE '2016-01-01' "
            "AND award_date < DATE '2026-01-01'"
        )
        total_awards, total_amount = connection.execute(
            f"SELECT count(*), coalesce(sum(contract_amount), 0) FROM philgeps WHERE {where}"
        ).fetchone()
        awardees = connection.execute(
            f"""
            SELECT awardee_name, count(*) AS awards, coalesce(sum(contract_amount), 0)
            FROM philgeps
            WHERE {where} AND awardee_name IS NOT NULL
            GROUP BY 1
            ORDER BY 3 DESC
            LIMIT 5
            """
        ).fetchall()
        top_award = connection.execute(
            f"""
            SELECT award_title, awardee_name, contract_amount, award_date
            FROM philgeps
            WHERE {where}
            ORDER BY contract_amount DESC
            LIMIT 1
            """
        ).fetchone()
    finally:
        connection.close()

    if not top_award:
        raise ValueError("No PhilGEPS awards matched San Fernando")
    return {
        "totalAwards": int(total_awards),
        "totalAmount": as_number(total_amount, "PhilGEPS total amount"),
        "topAwardees": [
            {
                "name": display_awardee(name),
                "awards": int(awards),
                "amount": as_number(amount, f"{name} PhilGEPS amount"),
            }
            for name, awards, amount in awardees
        ],
        "topAward": {
            "title": display_award_title(top_award[0]),
            "awardee": display_awardee(top_award[1]),
            "amount": as_number(top_award[2], "PhilGEPS top award amount"),
            "year": int(top_award[3].year),
        },
    }


def write_projects(records: list[dict[str, Any]], bounds: tuple[float, float, float, float]) -> None:
    min_lng, min_lat, max_lng, max_lat = bounds
    lines = [
        "export interface DpwhProjectRecord {",
        "  contractId: string",
        "  title: string",
        "  year: number",
        "  budget: number",
        "  status: string",
        "  progress: number",
        "  contractor: string",
        "  category: string",
        "  barangay: string | null",
        "  latitude: number | null",
        "  longitude: number | null",
        "  startDate: string | null",
        "  completionDate: string | null",
        "  sourceOfFunds: string | null",
        "}",
        "",
        "export const dpwhProjectSourceUrl = 'https://data.bettergov.ph/datasets/19'",
        "",
        "export const dpwhMapBounds = {",
        f"  minLat: {min_lat:.7f},",
        f"  maxLat: {max_lat:.7f},",
        f"  minLng: {min_lng:.7f},",
        f"  maxLng: {max_lng:.7f},",
        "}",
        "",
        "export const dpwhProjects: DpwhProjectRecord[] = [",
    ]
    for record in records:
        lines.extend(
            [
                "  {",
                f"    contractId: {quote(record['contractId'])},",
                f"    title: {quote(record['title'])},",
                f"    year: {record['year']},",
                f"    budget: {num(record['budget']) },",
                f"    status: {quote(record['status'])},",
                f"    progress: {num(record['progress'])},",
                f"    contractor: {quote(record['contractor'])},",
                f"    category: {quote(record['category'])},",
                f"    barangay: {quote(record['barangay']) if record['barangay'] is not None else 'null'},",
                f"    latitude: {num(record['latitude']) if record['latitude'] is not None else 'null'},",
                f"    longitude: {num(record['longitude']) if record['longitude'] is not None else 'null'},",
                f"    startDate: {quote(record['startDate']) if record['startDate'] is not None else 'null'},",
                f"    completionDate: {quote(record['completionDate']) if record['completionDate'] is not None else 'null'},",
                f"    sourceOfFunds: {quote(record['sourceOfFunds']) if record['sourceOfFunds'] is not None else 'null'},",
                "  },",
            ]
        )
    lines.extend([" ]", ""])
    write_text(DATA_DIR / "projects.ts", "\n".join(lines))


def generate_barangay_map(
    path: Path,
    detailed_city_path: Path | None = None,
) -> tuple[
    list[dict[str, Any]],
    str,
    tuple[float, float, float, float],
    str,
    list[dict[str, Any]],
    dict[str, float],
]:
    source_text = (DATA_DIR / "barangays.ts").read_text(encoding="utf-8")
    declaration = source_text.index("export const barangays")
    array_start = source_text.index("[", source_text.index("= [", declaration))
    array_end = source_text.rindex("]")
    sites = json.loads(source_text[array_start : array_end + 1])
    site_by_key = {place_key(site["name"]): site for site in sites}

    with path.open(encoding="utf-8") as handle:
        geojson = json.load(handle)
    city_features = [
        feature
        for feature in geojson["features"]
        if feature["properties"].get("ADM3_EN") == "City of San Fernando (Capital)"
        and feature["properties"].get("ADM2_EN") == "La Union"
    ]
    features = city_features
    if detailed_city_path is not None and detailed_city_path.is_file():
        with detailed_city_path.open(encoding="utf-8") as handle:
            detailed_geojson = json.load(handle)
        features = detailed_geojson["features"]
    if not features:
        raise ValueError("No San Fernando features found in the boundary GeoJSON")

    geometries = [shape(feature["geometry"]) for feature in features]
    bounds = [geometry.bounds for geometry in geometries]
    minx = min(bound[0] for bound in bounds)
    maxx = max(bound[2] for bound in bounds)
    miny = min(bound[1] for bound in bounds)
    maxy = max(bound[3] for bound in bounds)
    earth_radius = 6_378_137.0

    def project(x: float, y: float, z: Any = None) -> tuple[float, float]:
        latitude = math.radians(y)
        return (
            earth_radius * math.radians(x),
            -earth_radius * math.log(math.tan(math.pi / 4 + latitude / 2)),
        )

    def project_coords(coords: Any) -> Any:
        projected = coords.copy()
        latitudes = [math.radians(latitude) for latitude in coords[:, 1]]
        projected[:, 0] = [earth_radius * math.radians(longitude) for longitude in coords[:, 0]]
        projected[:, 1] = [
            -earth_radius * math.log(math.tan(math.pi / 4 + latitude / 2))
            for latitude in latitudes
        ]
        return projected

    map_minx = minx - 0.012
    map_miny = miny - 0.010
    map_maxx = maxx + 0.012
    map_maxy = maxy + 0.010
    x0, y0 = project(map_minx, map_maxy)
    x1, y1 = project(map_maxx, map_miny)
    width = x1 - x0
    height = y1 - y0
    view_width = 1000.0
    view_height = view_width * height / width
    padding = 12.0
    scale_x = (view_width - padding * 2) / width
    scale_y = (view_height - padding * 2) / height

    def geometry_path(
        geometry: Any,
        include_interiors: bool = True,
        already_projected: bool = False,
    ) -> str:
        projected = (
            geometry
            if already_projected
            else transform_geometry(geometry, project_coords)
        ).simplify(24, preserve_topology=True)
        polygons = (
            list(projected.geoms) if projected.geom_type == "MultiPolygon" else [projected]
        )
        parts: list[str] = []
        for polygon in polygons:
            rings = [polygon.exterior, *polygon.interiors] if include_interiors else [polygon.exterior]
            for ring in rings:
                coordinates = [
                    ((x - x0) * scale_x + padding, (y - y0) * scale_y + padding)
                    for x, y in ring.coords
                ]
                if len(coordinates) < 4:
                    continue
                parts.append(
                    "M"
                    + " ".join(f"{x:.1f},{y:.1f}" for x, y in coordinates[:-1])
                    + "Z"
                )
        return "".join(parts)

    def map_point(x: float, y: float) -> tuple[float, float]:
        return ((x - x0) * scale_x + padding, (y - y0) * scale_y + padding)

    tile_zoom = MAP_TILE_ZOOM
    tile_count = 2**tile_zoom
    world_half = math.pi * earth_radius
    tile_span = (2 * world_half) / tile_count
    min_tile_x = max(0, math.floor((x0 + world_half) / tile_span))
    max_tile_x = min(tile_count - 1, math.floor((x1 + world_half) / tile_span))
    min_tile_y = max(0, math.floor((y0 + world_half) / tile_span))
    max_tile_y = min(tile_count - 1, math.floor((y1 + world_half) / tile_span))
    tiles = []
    for tile_y in range(min_tile_y, max_tile_y + 1):
        for tile_x in range(min_tile_x, max_tile_x + 1):
            left = -world_half + tile_x * tile_span
            top = -world_half + tile_y * tile_span
            right = left + tile_span
            bottom = top + tile_span
            screen_left, screen_top = map_point(left, top)
            screen_right, screen_bottom = map_point(right, bottom)
            tiles.append(
                {
                    "tileZ": tile_zoom,
                    "tileX": tile_x,
                    "tileY": tile_y,
                    "x": screen_left,
                    "y": screen_top,
                    "width": screen_right - screen_left,
                    "height": screen_bottom - screen_top,
                }
            )

    items: list[dict[str, Any]] = []
    unmatched: list[str] = []
    used_slugs: set[str] = set()
    for feature, geometry in zip(features, geometries):
        properties = feature["properties"]
        geo_name = properties["ADM4_EN"]
        key = place_key(geo_name)
        if key not in site_by_key and geo_name == "Cabaroan (Negro)":
            key = "cabaroan"
        site = site_by_key.get(key)
        if not site:
            unmatched.append(geo_name)
            continue
        used_slugs.add(site["slug"])
        projected_geometry = transform_geometry(geometry, project_coords)
        label_x, label_y = map_point(
            *projected_geometry.representative_point().coords[0]
        )
        items.append(
            {
                "slug": site["slug"],
                "name": site["name"],
                "areaSqKm": as_number(properties.get("AREA_SQKM") or 0, f"{geo_name} area"),
                "labelX": label_x,
                "labelY": label_y,
                "path": geometry_path(projected_geometry, already_projected=True),
            }
        )

    if unmatched:
        raise ValueError(f"Unmatched boundary features: {', '.join(sorted(unmatched))}")
    missing = sorted(
        site["name"] for site in sites if site["slug"] not in used_slugs
    )
    if missing:
        raise ValueError(f"Barangays without boundary features: {', '.join(missing)}")

    city_outline = geometry_path(
        unary_union(
            [transform_geometry(geometry, project_coords) for geometry in geometries]
        ),
        include_interiors=False,
        already_projected=True,
    )

    items.sort(key=lambda item: item["name"])
    return (
        items,
        f"0 0 {view_width:.0f} {view_height:.0f}",
        (minx, miny, maxx, maxy),
        city_outline,
        tiles,
        {
            "originX": x0,
            "originY": y0,
            "scaleX": scale_x,
            "scaleY": scale_y,
            "padding": padding,
        },
    )


def write_bytes(path: Path, data: bytes) -> None:
    with tempfile.NamedTemporaryFile(
        mode="wb",
        dir=path.parent,
        prefix=f"{path.name}.",
        suffix=".tmp",
        delete=False,
    ) as handle:
        handle.write(data)
        temp_path = Path(handle.name)
    os.replace(temp_path, path)


def write_text(path: Path, text: str) -> None:
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        newline="\n",
        dir=path.parent,
        prefix=f"{path.name}.",
        suffix=".tmp",
        delete=False,
    ) as handle:
        handle.write(text)
        temp_path = Path(handle.name)
    os.replace(temp_path, path)


def write_fiscal(fiscal: dict[str, Any], peers: list[dict[str, Any]]) -> None:
    lines = [
        "export interface RevenueSource {",
        "  key: 'tax' | 'nonTax' | 'nta' | 'otherNational' | 'grants'",
        "  label: string",
        "  amount: number",
        "}",
        "",
        "export interface CityFiscalPeer {",
        "  name: string",
        "  province: string",
        "  annualRegularIncome: number",
        "  localSources: number",
        "  ownSourceShare: number",
        "  ntaShare: number",
        "}",
        "",
        "export const fiscalSourceUrl = 'https://data.bettergov.ph/datasets/9'",
        "",
        "export const fiscal2024 = {",
        f"  totalIncome: {num(fiscal['totalIncome'])},",
        f"  totalLocalSources: {num(fiscal['totalLocalSources'])},",
        f"  totalTaxRevenue: {num(fiscal['totalTaxRevenue'])},",
        f"  totalNonTaxRevenue: {num(fiscal['totalNonTaxRevenue'])},",
        f"  nta: {num(fiscal['nta'])},",
        f"  otherNationalShares: {num(fiscal['otherNationalShares'])},",
        f"  grants: {num(fiscal['grants'])},",
        f"  totalOperatingExpenditures: {num(fiscal['totalOperatingExpenditures'])},",
        f"  netOperatingIncome: {num(fiscal['netOperatingIncome'])},",
        f"  capitalOutlay: {num(fiscal['capitalOutlay'])},",
        f"  specialEducationFund: {num(fiscal['specialEducationFund'])},",
        "  disasterFund: {",
        f"    appropriation: {num(fiscal['disasterFund']['appropriation'])},",
        f"    spent: {num(fiscal['disasterFund']['spent'])},",
        f"    quickResponseAppropriation: {num(fiscal['disasterFund']['quickResponseAppropriation'])},",
        f"    quickResponseSpent: {num(fiscal['disasterFund']['quickResponseSpent'])},",
        "  },",
        "}",
        "",
        "export const revenueSources2024: RevenueSource[] = [",
        "  { key: 'tax', label: 'Local taxes', amount: fiscal2024.totalTaxRevenue },",
        "  { key: 'nonTax', label: 'Fees and other local revenue', amount: fiscal2024.totalNonTaxRevenue },",
        "  { key: 'nta', label: 'National Tax Allotment', amount: fiscal2024.nta },",
        "  { key: 'otherNational', label: 'Other national shares', amount: fiscal2024.otherNationalShares },",
        "  { key: 'grants', label: 'Grants and extraordinary receipts', amount: fiscal2024.grants },",
        "]",
        "",
        "export const ilocosCities2024: CityFiscalPeer[] = [",
    ]
    for peer in peers:
        lines.extend(
            [
                "  {",
                f"    name: {quote(peer['name'])},",
                f"    province: {quote(peer['province'])},",
                f"    annualRegularIncome: {num(peer['annualRegularIncome'])},",
                f"    localSources: {num(peer['localSources'])},",
                f"    ownSourceShare: {num(peer['ownSourceShare'], 4)},",
                f"    ntaShare: {num(peer['ntaShare'], 4)},",
                "  },",
            ]
        )
    lines.extend(["]", ""])
    write_text(DATA_DIR / "fiscal.ts", "\n".join(lines))


def write_elections(elections: list[dict[str, Any]]) -> None:
    lines = [
        "export interface ElectionWinner {",
        "  name: string",
        "  party: string",
        "}",
        "",
        "export interface CityElection {",
        "  year: number",
        "  mayor: ElectionWinner | null",
        "  viceMayor: ElectionWinner | null",
        "}",
        "",
        "export interface ElectionSupplementSource {",
        "  year: number",
        "  url: string",
        "}",
        "",
        "export const electionSourceUrl = 'https://data.bettergov.ph/datasets/25'",
        "export const electionSupplementSources: ElectionSupplementSource[] = [",
    ]
    for year, supplement in sorted(ELECTION_SUPPLEMENTS.items()):
        lines.append(
            "  { "
            f"year: {year}, url: {quote(supplement['sourceUrl'])} "
            "},"
        )
    lines.extend(
        [
            "]",
            "",
            "export const cityElections: CityElection[] = [",
        ]
    )
    for election in elections:
        lines.extend(["  {", f"    year: {election['year']},"])
        for key in ("mayor", "viceMayor"):
            winner = election[key]
            if winner is None:
                lines.append(f"    {key}: null,")
            else:
                lines.append(
                    f"    {key}: {{ name: {quote(winner['name'])}, party: {quote(winner['party'])} }},"
                )
        lines.append("  },")
    lines.extend(["]", ""])
    write_text(DATA_DIR / "elections.ts", "\n".join(lines))


def write_contracts(dpwh: dict[str, Any], philgeps: dict[str, Any]) -> None:
    lines = [
        "export interface DpwhProject {",
        "  contractId: string",
        "  title: string",
        "  year: number",
        "  budget: number",
        "  status: string",
        "  progress: number",
        "  contractor: string",
        "}",
        "",
        "export interface YearTotal {",
        "  year: number",
        "  projects: number",
        "  amount: number",
        "}",
        "",
        "export interface AwardeeTotal {",
        "  name: string",
        "  awards: number",
        "  amount: number",
        "}",
        "",
        "export const dpwhSourceUrl = 'https://data.bettergov.ph/datasets/19'",
        "export const philgepsSourceUrl = 'https://data.bettergov.ph/datasets/5'",
        "",
        "export const dpwhSanFernando = {",
        f"  totalProjects: {dpwh['totalProjects']},",
        f"  totalBudget: {num(dpwh['totalBudget'])},",
        f"  completed: {dpwh['completed']},",
        f"  ongoing: {dpwh['ongoing']},",
        f"  forProcurement: {dpwh['forProcurement']},",
        "  recentYears: [",
    ]
    for year in dpwh["recentYears"]:
        lines.append(
            "    { "
            f"year: {year['year']}, projects: {year['projects']}, "
            f"amount: {num(year['amount'])} "
            "},"
        )
    lines.extend(["  ] satisfies YearTotal[],", "  largestProjects: ["])
    for project in dpwh["largestProjects"]:
        lines.extend(
            [
                "    {",
                f"      contractId: {quote(project['contractId'])},",
                f"      title: {quote(project['title'])},",
                f"      year: {project['year']},",
                f"      budget: {num(project['budget'])},",
                f"      status: {quote(project['status'])},",
                f"      progress: {num(project['progress'])},",
                f"      contractor: {quote(project['contractor'])},",
                "    },",
            ]
        )
    lines.extend(
        [
            "  ] satisfies DpwhProject[],",
            "}",
            "",
            "export const philgepsSanFernando = {",
            f"  totalAwards: {philgeps['totalAwards']},",
            f"  totalAmount: {num(philgeps['totalAmount'])},",
            "  topAwardees: [",
        ]
    )
    for awardee in philgeps["topAwardees"]:
        lines.append(
            "    { "
            f"name: {quote(awardee['name'])}, awards: {awardee['awards']}, "
            f"amount: {num(awardee['amount'])} "
            "},"
        )
    lines.extend(
        [
            "  ] satisfies AwardeeTotal[],",
            "  topAward: {",
            f"    title: {quote(philgeps['topAward']['title'])},",
            f"    awardee: {quote(philgeps['topAward']['awardee'])},",
            f"    amount: {num(philgeps['topAward']['amount'])},",
            f"    year: {philgeps['topAward']['year']},",
            "  },",
            "}",
            "",
        ]
    )
    write_text(DATA_DIR / "contracts.ts", "\n".join(lines))


def write_barangay_map(
    features: list[dict[str, Any]],
    view_box: str,
    city_outline: str,
    tiles: list[dict[str, Any]],
    projection: dict[str, float],
) -> None:
    lines = [
        "export interface BarangayMapFeature {",
        "  slug: string",
        "  name: string",
        "  areaSqKm: number",
        "  labelX: number",
        "  labelY: number",
        "  path: string",
        "}",
        "",
        "export interface BarangayMapProjection {",
        "  originX: number",
        "  originY: number",
        "  scaleX: number",
        "  scaleY: number",
        "  padding: number",
        "}",
        "",
        "export interface BarangayMapTile {",
        "  tileZ: number",
        "  tileX: number",
        "  tileY: number",
        "  x: number",
        "  y: number",
        "  width: number",
        "  height: number",
        "}",
        "",
        "export const barangayMap = {",
        f"  viewBox: {quote(view_box)},",
        f"  cityOutline: {quote(city_outline)},",
        "  projection: {",
        f"    originX: {projection['originX']:.12f},",
        f"    originY: {projection['originY']:.12f},",
        f"    scaleX: {projection['scaleX']:.12f},",
        f"    scaleY: {projection['scaleY']:.12f},",
        f"    padding: {projection['padding']:.1f},",
        "  },",
        "  tiles: [",
    ]
    for tile in tiles:
        lines.extend(
            [
                "    {",
                f"      tileZ: {int(tile['tileZ'])},",
                f"      tileX: {int(tile['tileX'])},",
                f"      tileY: {int(tile['tileY'])},",
                f"      x: {tile['x']:.2f},",
                f"      y: {tile['y']:.2f},",
                f"      width: {tile['width']:.2f},",
                f"      height: {tile['height']:.2f},",
                "    },",
            ]
        )
    lines.extend(["  ],", "  features: ["])
    for feature in features:
        lines.extend(
            [
                "    {",
                f"      slug: {quote(feature['slug'])},",
                f"      name: {json.dumps(feature['name'], ensure_ascii=False)},",
                f"      areaSqKm: {feature['areaSqKm']:.2f},",
                f"      labelX: {feature['labelX']:.1f},",
                f"      labelY: {feature['labelY']:.1f},",
                f"      path: {quote(feature['path'])},",
                "    },",
            ]
        )
    lines.extend(["  ],", "} as const", ""])
    write_text(DATA_DIR / "barangayMap.ts", "\n".join(lines))


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Refresh the checked-in public data snapshots."
    )
    parser.add_argument(
        "--cache-dir",
        type=Path,
        default=DEFAULT_CACHE_DIR,
        help="Directory for downloaded source files.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Download every source again even if the cached copy is current.",
    )
    parser.add_argument(
        "--offline",
        action="store_true",
        help="Use only cached source files; do not contact the API or download.",
    )
    parser.add_argument(
        "--news-only",
        action="store_true",
        help="Refresh only the news snapshot and images.",
    )
    args = parser.parse_args()
    if args.force and args.offline:
        parser.error("--force and --offline cannot be used together")
    cached_fallbacks: list[str] = []
    if args.news_only:
        news_count = refresh_news(args.offline, cached_fallbacks)
        write_refresh_status(cached_fallbacks)
        print(f"refreshed {news_count or 0} news stories")
        return
    cache_dir = args.cache_dir.resolve()
    cache_dir.mkdir(parents=True, exist_ok=True)

    sources = resolve_sources(args.offline)
    paths = {
        key: download_source(key, spec, cache_dir, args.force, args.offline, cached_fallbacks)
        for key, spec in sources.items()
    }

    print("parsing sources")
    fiscal, peers = parse_fiscal(
        paths["sre"], paths["ari"], paths["ldrrmf"], paths["sef"]
    )
    elections = parse_elections(paths["elections"])
    dpwh = parse_dpwh(paths["dpwh"])
    philgeps = parse_philgeps(paths["philgeps"])
    map_features, view_box, map_bounds, city_outline, map_tiles, map_projection = generate_barangay_map(
        paths["boundaries"], CITY_BOUNDARIES_PATH
    )

    print("writing snapshots")
    write_fiscal(fiscal, peers)
    write_elections(elections)
    write_contracts(dpwh, philgeps)
    write_projects(dpwh["records"], map_bounds)
    write_barangay_map(map_features, view_box, city_outline, map_tiles, map_projection)
    news_count = refresh_news(args.offline, cached_fallbacks)

    summary = (
        "refreshed "
        f"{len(peers)} cities, {len(elections)} election cycles, "
        f"{dpwh['totalProjects']} DPWH projects, "
        f"{philgeps['totalAwards']} PhilGEPS awards, "
        f"{len(map_features)} barangay boundaries, "
        f"{len(dpwh['records'])} project records"
    )
    if news_count is not None:
        summary += f", {news_count} news stories"
    print(summary)
    write_refresh_status(cached_fallbacks)


def write_refresh_status(cached_fallbacks: list[str]) -> None:
    cached_fallbacks = list(dict.fromkeys(cached_fallbacks))
    output_path = os.environ.get("GITHUB_OUTPUT")
    if output_path:
        with open(output_path, "a", encoding="utf-8") as output:
            output.write(f"cached_fallbacks={','.join(cached_fallbacks)}\n")

    if not cached_fallbacks:
        return

    names = ", ".join(f"`{name}`" for name in cached_fallbacks)
    print(f"warning: refresh used cached data after source failures: {', '.join(cached_fallbacks)}")
    summary_path = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary_path:
        with open(summary_path, "a", encoding="utf-8") as summary:
            summary.write(
                "### Cached source fallback\n\n"
                f"The refresh reused cached data for {names}. No refresh PR will be opened for this run.\n"
            )
    if os.environ.get("GITHUB_ACTIONS") == "true":
        for name in cached_fallbacks:
            print(f"::warning title=Cached source used::{name} used its cached copy after a refresh failure")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"error: {error}", file=sys.stderr)
        raise
